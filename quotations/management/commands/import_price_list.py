import openpyxl

from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

from quotations.models import Product


PRICE_LIST_COLUMNS = {
    1: "acdb",
    3: "dcdb",
    5: "ac_wire",
    7: "dc_cable",
    9: "energy_meter",
    11: "meter_box",
    13: "ug_cable",
    15: "net_meter",
}


def clean(value):
    if value is None:
        return ""

    return str(value).strip()


def decimal_value(value):
    if value is None or value == "":
        return Decimal("0")

    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0")


class Command(BaseCommand):

    help = "Import PRICE LIST specifications and rates into Product."

    def add_arguments(self, parser):
        parser.add_argument(
            "excel_file",
            type=str,
            help="Path to the quotation Excel workbook",
        )

        parser.add_argument(
            "--clear",
            action="store_true",
            help="Clear existing Product records before importing.",
        )

    def handle(self, *args, **options):

        excel_file = Path(options["excel_file"])

        if not excel_file.exists():
            raise CommandError(
                f"Excel file not found: {excel_file}"
            )

        self.stdout.write(
            self.style.NOTICE(
                f"Reading workbook: {excel_file}"
            )
        )

        workbook = openpyxl.load_workbook(
            excel_file,
            read_only=True,
            data_only=True,
        )

        sheet_name = "PRICE LIST "

        if sheet_name not in workbook.sheetnames:
            raise CommandError(
                f"Sheet '{sheet_name}' was not found."
            )

        worksheet = workbook[sheet_name]

        if options["clear"]:
            Product.objects.all().delete()

            self.stdout.write(
                self.style.WARNING(
                    "Existing Product records cleared."
                )
            )

        imported = 0
        skipped = 0

        # Rows 1 and 2 contain category/specification headers.
        # Actual price-list records start from row 3.
        for row_number in range(3, worksheet.max_row + 1):

            for column_number, category in PRICE_LIST_COLUMNS.items():

                specification = clean(
                    worksheet.cell(
                        row=row_number,
                        column=column_number
                    ).value
                )

                rate = worksheet.cell(
                    row=row_number,
                    column=column_number + 1
                ).value

                if not specification:
                    continue

                if rate is None or rate == "":
                    skipped += 1
                    continue

                Product.objects.create(
                    name=specification,
                    category=category,
                    make="",
                    specification=specification,
                    unit="Nos",
                    rate=decimal_value(rate),
                    tax_percent=18,
                    is_active=True,
                )

                imported += 1

        self.stdout.write("")

        self.stdout.write(
            self.style.SUCCESS(
                "PRICE LIST import completed successfully."
            )
        )

        self.stdout.write(
            f"Products imported: {imported}"
        )

        self.stdout.write(
            f"Rows skipped: {skipped}"
        )

        self.stdout.write(
            f"Total products in database: {Product.objects.count()}"
        )