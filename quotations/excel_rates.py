"""
Reads the Dehlsen rate sheet (QOUTATION_WORK_OUT.xlsx) and turns it into plain Python rows.

Pure Python + openpyxl: no Django, so it can be tested on its own.  The management command
`load_rates_from_excel` writes these rows into the admin tables.

The sheet is read by its SECTION TITLES, not by fixed row numbers, so adding rows to a section later
does not break the reader.  The BOS section is ignored on purpose.  The 'Non Taxable Particular'
section is read but reported only (those amounts are fields on the quotation, not tables).
"""
import re
from decimal import Decimal

import openpyxl

# section title (upper-case, single spaces)  ->  key
SECTIONS = {
    "SOLAR PANEL MOD": "panels",
    "POLY CAB INVERTER INV": "inverters",
    "ACDB DCDB BOX": "boxes",
    "METER BOX": "meter_box",
    "DC CABLE": "dc_cable",
    "AC CABLE": "ac_cable",
    "AC CABLE UG": "ug_cable",
    "EARTHING CABLE": "earthing_cable",
    "EARTH ROD AND LIGHTING ARRESTOR": "earth_la",
    "METERS": "meters",
    "BOS": "bos",
    "NON TAXABLE PARTICULAR": "charges",
}


def plain(value):
    """Decimal -> plain text without scientific notation:  Decimal('10') -> '10',  8000 -> '8000',  4.50 -> '4.5'."""
    return format(Decimal(value).normalize(), "f")


def _norm(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().upper()


def _text(value):
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def _dec(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value).strip())
    except Exception:
        return None


def _kw(value):
    """'3KW' -> Decimal('3')."""
    m = re.search(r"(\d+(?:\.\d+)?)", _text(value))
    return Decimal(m.group(1)) if m else None


def _phase(value):
    """'1PH', '1 PH', 1, '3PH', 3, 'SINGLE', 'THREE'  ->  '1PH' / '3PH'."""
    t = _norm(value).replace(" ", "")
    if t in ("1", "1PH", "SINGLE", "SINGLEPHASE"):
        return "1PH"
    if t in ("3", "3PH", "THREE", "THREEPHASE"):
        return "3PH"
    return _text(value)


def _find_sections(ws):
    """-> list of (key, title_row) in sheet order."""
    found = []
    for row in ws.iter_rows(min_col=1, max_col=2):
        for cell in row:
            key = SECTIONS.get(_norm(cell.value))
            if key and cell.value is not None:
                found.append((key, cell.row))
                break
    return found


def _section_rows(ws, start, end):
    """Data rows between a title and the next title: skips the header row ('SL NO') and blank rows."""
    for r in range(start + 1, end):
        values = [ws.cell(row=r, column=c).value for c in range(1, 8)]
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        if _norm(values[0]) == "SL NO":
            continue
        yield r, values


def parse_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    sections = _find_sections(ws)
    bounds = {}
    for i, (key, row) in enumerate(sections):
        bounds[key] = (row, sections[i + 1][1] if i + 1 < len(sections) else ws.max_row + 1)

    out = {"panels": [], "inverters": [], "products": [], "charges": [], "skipped": [], "bos_ignored_rows": 0}
    warnings = out["skipped"]

    def rows(key):
        return list(_section_rows(ws, *bounds[key])) if key in bounds else []

    # ---- panels: B company, C technology, D watts, E rate per watt --------------------------
    for r, v in rows("panels"):
        company, tech, watts, rate = _text(v[1]), _text(v[2]), _dec(v[3]), _dec(v[4])
        if not company or watts is None or rate is None:
            warnings.append("Panels row %d skipped (needs company, watts and rate/watt)" % r)
            continue
        out["panels"].append({"company": company, "panel_type": tech, "wattage": watts, "rate_per_watt": rate, "row": r})

    # ---- inverters: B brand, C phase, D mppt, E kw, F price ---------------------------------
    for r, v in rows("inverters"):
        brand, phase, mppt, kw, price = _text(v[1]), _phase(v[2]), _text(v[3]).upper(), _kw(v[4]), _dec(v[5])
        if not brand or kw is None or price is None:
            warnings.append("Inverter row %d skipped" % r)
            continue
        out["products"].append({
            "category": "inverter", "make": brand, "unit": "Nos", "rate": price, "inverter_capacity": kw,
            "name": "%s %sKW %s %s" % (brand, plain(kw), phase, mppt),
            "specification": "%s KW %s on-grid inverter, %s MPPT" % (plain(kw), "SINGLE phase" if phase == "1PH" else "THREE phase", mppt),
            "row": r})

    # ---- ACDB / DCDB: B box, C spec, D phase, E kw range or voltage, F price -----------------
    for r, v in rows("boxes"):
        box, spec, phase, extra, price = _norm(v[1]), _text(v[2]).upper(), _phase(v[3]), _text(v[4]), _dec(v[5])
        if box not in ("ACDB", "DCDB") or price is None:
            warnings.append("ACDB/DCDB row %d skipped" % r)
            continue
        out["products"].append({
            "category": box.lower(), "make": "", "unit": "Nos", "rate": price,
            "name": "%s %s %s" % (box, spec, phase), "specification": "%s %s, %s, %s" % (box, spec, phase, extra), "row": r})

    # ---- meter box: B box, C phase, D kw range, E price --------------------------------------
    for r, v in rows("meter_box"):
        box, phase, extra, price = _norm(v[1]), _phase(v[2]), _text(v[3]), _dec(v[4])
        if price is None:
            warnings.append("Meter box row %d skipped" % r)
            continue
        out["products"].append({
            "category": "meter_box", "make": "", "unit": "Nos", "rate": price,
            "name": "METER BOX %s %s" % (box, phase), "specification": "Meter box, %s, %s" % (phase, extra), "row": r})

    # ---- cables: B type, C brand, D size / spec, E price per metre ------------------------------
    for key, category in (("dc_cable", "dc_cable"), ("ac_cable", "ac_wire"), ("ug_cable", "ug_cable")):
        for r, v in rows(key):
            brand, size, price = _text(v[2]), _text(v[3]), _dec(v[4])
            if not size or price is None:
                warnings.append("%s row %d skipped" % (key, r))
                continue
            if key == "dc_cable":
                name, spec = "%s SQMM" % size, "DC cable %s sq.mm" % size                    # matches the page's '4 SQMM' options
            elif key == "ac_cable":
                name, spec = "%s SQMM FR Copper - %s" % (size, brand), "AC wire %s sq.mm FR copper" % size
            else:
                name, spec = "UG CABLE %s - %s" % (size, brand), "AC UG cable %s" % size
            out["products"].append({"category": category, "make": brand, "unit": "m", "rate": price,
                                    "name": name, "specification": spec, "row": r})

    for r, v in rows("earthing_cable"):
        brand, spec, price = _text(v[2]), _text(v[3]), _dec(v[4])
        if price is None:
            warnings.append("Earthing cable row %d skipped" % r)
            continue
        out["products"].append({"category": "earthing", "make": brand, "unit": "m", "rate": price,
                                "name": "EARTHING CABLE %s - %s" % (spec, brand), "specification": "Earthing cable %s" % spec, "row": r})

    # ---- earth rod / lightning arrestor: B name, C brand, D spec, E price -----------------------
    for r, v in rows("earth_la"):
        name, brand, spec, price = _norm(v[1]), _text(v[2]), _text(v[3]).upper(), _dec(v[4])
        if price is None:
            warnings.append("Earth rod / arrestor row %d skipped" % r)
            continue
        if "ROD" in name:
            product = ("earthing", "EARTH ROD")
        else:
            product = ("electrical", "LIGHTNING ARRESTOR")          # the sheet spells it 'LIGTINING'
        out["products"].append({"category": product[0], "make": brand, "unit": "Nos", "rate": price,
                                "name": "%s %s" % (product[1], spec), "specification": "%s %s" % (product[1].title(), spec.title()), "row": r})

    # ---- meters: B name, C brand, D phase, E price --------------------------------------------------
    for r, v in rows("meters"):
        name, brand, phase, price = _norm(v[1]), _text(v[2]), _phase(v[3]), _dec(v[4])
        category = "energy_meter" if "ENERGY" in name else "net_meter" if "NET" in name else None
        if category is None or price is None:
            warnings.append("Meter row %d skipped" % r)
            continue
        out["products"].append({"category": category, "make": brand, "unit": "Nos", "rate": price,
                                "name": "%s %s" % (name, phase), "specification": "%s, %s, class 1, LT TOD" % (name.title(), phase), "row": r})

    # ---- charges: reported only ----------------------------------------------------------------------
    for r, v in rows("charges"):
        item, rate = _text(v[1]), _dec(v[2])
        if item and rate is not None:
            out["charges"].append({"item": item, "rate": rate})

    out["bos_ignored_rows"] = len(rows("bos"))

    # ---- product names must be unique inside a category: add the brand when two rows would clash --------
    seen = {}
    for p in out["products"]:
        key = (p["category"], p["name"].upper())
        if key in seen:
            p["name"] = "%s - %s" % (p["name"], p["make"] or "row %d" % p["row"])
        seen[(p["category"], p["name"].upper())] = True

    return out
 